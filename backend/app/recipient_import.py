from __future__ import annotations

import csv
import io
from pathlib import Path

HEADER_ALIASES = {
    "target", "recipient", "receiver", "username", "user_name", "user name",
    "usname", "telegram", "telegram_username", "telegram username",
    "phone", "phone_number", "phone number", "mobile", "telephone",
    "số điện thoại", "so dien thoai", "sdt",
}

def _clean_header(value) -> str:
    return str(value or "").strip().lower().replace("-", "_")

def _stringify(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def load_recipient_rows(filename: str, data: bytes) -> list[list[object]]:
    suffix = Path(filename or "").suffix.lower()
    if suffix == ".csv":
        text = None
        for enc in ("utf-8-sig", "utf-8", "cp1258", "latin-1"):
            try:
                text = data.decode(enc); break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise ValueError("Không thể đọc mã hóa của tệp CSV")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = [list(r) for r in csv.reader(io.StringIO(text), dialect)]
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            rows = [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    elif suffix == ".xls":
        import xlrd
        book = xlrd.open_workbook(file_contents=data)
        sheet = book.sheet_by_index(0)
        rows = [sheet.row_values(i) for i in range(sheet.nrows)]
    else:
        raise ValueError("Chỉ hỗ trợ tệp .csv, .xlsx hoặc .xls")
    return [list(r) for r in rows if any(_stringify(v) for v in r)]


def preview_recipient_file(filename: str, data: bytes) -> dict:
    rows = load_recipient_rows(filename, data)
    if not rows:
        raise ValueError("Tệp không có dữ liệu người nhận")
    width = min(50, max(len(r) for r in rows))
    first = [_stringify(rows[0][i]) if i < len(rows[0]) else "" for i in range(width)]
    suggested = [i for i, value in enumerate(first) if _clean_header(value) in HEADER_ALIASES]
    header_detected = bool(suggested)
    headers = [value or f"Cột {i + 1}" for i, value in enumerate(first)] if header_detected else [f"Cột {i + 1}" for i in range(width)]
    sample_rows = rows[1:9] if header_detected else rows[:8]
    sample = [[_stringify(row[i]) if i < len(row) else "" for i in range(width)] for row in sample_rows]
    return {
        "headers": headers, "suggested_columns": suggested,
        "header_detected": header_detected, "rows": len(rows) - (1 if header_detected else 0),
        "sample": sample, "column_count": width,
    }

def extract_recipient_columns(filename: str, data: bytes, selected: list[int] | None = None, has_header: bool | None = None) -> dict:
    rows = load_recipient_rows(filename, data)
    if not rows:
        raise ValueError("Tệp không có dữ liệu người nhận")
    width = max(len(r) for r in rows)
    first = [_stringify(rows[0][i]) if i < len(rows[0]) else "" for i in range(width)]
    auto = [i for i, value in enumerate(first) if _clean_header(value) in HEADER_ALIASES]
    if selected is None:
        selected = auto
        if not selected:
            if width != 1:
                raise ValueError("Không xác định được cột người nhận. Hãy dùng chế độ xem trước và chọn cột.")
            selected = [0]
    selected = list(dict.fromkeys(int(i) for i in selected))
    if not selected or any(i < 0 or i >= width for i in selected):
        raise ValueError("Cột người nhận được chọn không hợp lệ")
    header = bool(auto) if has_header is None else bool(has_header)
    data_rows = rows[1:] if header else rows
    out: list[str] = []
    for row in data_rows:
        for idx in selected:
            value = _stringify(row[idx]) if idx < len(row) else ""
            if value:
                out.append(value)
    if not out:
        raise ValueError("Không tìm thấy username hoặc số điện thoại trong cột đã chọn")
    columns = [(first[i] or f"Cột {i + 1}") if header else f"Cột {i + 1}" for i in selected]
    return {"raw_targets": out, "columns": columns, "rows": len(data_rows)}

def parse_recipient_file(filename: str, data: bytes) -> dict:
    return extract_recipient_columns(filename, data)
