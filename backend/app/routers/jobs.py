import csv
import io
from collections import Counter

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from .. import job_store
from ..db import get_db
from ..models import Account, BulkJob, BulkJobItem, MessageDispatchItem, PhoneCheckItem
from ..audit import log_audit, _sanitize
from ..utils import bulk_stream, friendly_error
from ..realtime_events import emit_event
from ..job_timeline import build_job_timeline

router = APIRouter(prefix="/api/jobs", tags=["jobs"])

RETRYABLE_JOB_TYPES = {
    "group_join",
    "group_leave",
    "group_leave_target",
    "message_view",
    "terminate_other_sessions",
}
RETRY_ITEM_STATUSES = {"failed", "pending", "queued", "running"}


def _job_dict(job: BulkJob) -> dict:
    return {
        "id": job.id,
        "type": job.type,
        "status": job.status,
        "parameters": _sanitize(job.parameters or {}) if job.parameters else None,
        "retry_supported": job.type in RETRYABLE_JOB_TYPES,
        "total": job.total,
        "success": job.success,
        "failed": job.failed,
        "skipped": job.skipped,
        "pending": job.pending,
        "created_at": job.created_at,
        "started_at": job.started_at,
        "finished_at": job.finished_at,
        "heartbeat_at": job.heartbeat_at,
        "updated_at": job.updated_at,
        "paused_at": job.paused_at,
        "resume_count": job.resume_count or 0,
        "checkpoint": job.checkpoint or {},
        "last_error_code": job.last_error_code,
        "last_error_detail": job.last_error_detail,
        "resumable": job.type in job_store.RESUMABLE_JOB_TYPES,
    }


@router.get("")
async def list_jobs(
    limit: int = Query(50, ge=1, le=200),
    db: AsyncSession = Depends(get_db),
):
    res = await db.execute(
        select(BulkJob).order_by(BulkJob.created_at.desc()).limit(limit)
    )
    return [_job_dict(job) for job in res.scalars().all()]


@router.get("/{job_id}")
async def get_job(job_id: str, db: AsyncSession = Depends(get_db)):
    job = await db.get(BulkJob, job_id)
    if not job:
        raise HTTPException(404, "Không tìm thấy tác vụ")
    if job.type == "message_multi_send":
        res = await db.execute(
            select(MessageDispatchItem)
            .where(MessageDispatchItem.job_id == job_id)
            .order_by(MessageDispatchItem.id)
        )
        items = [{
            "id": item.id,
            "account_id": item.account_id,
            "target": item.target,
            "status": item.status,
            "attempts": item.attempts,
            "error_code": item.error_code,
            "error_detail": item.detail,
            "started_at": item.started_at,
            "finished_at": item.finished_at,
            "next_retry_at": item.next_retry_at,
        } for item in res.scalars().all()]
    elif job.type == "phone_check":
        res = await db.execute(
            select(PhoneCheckItem)
            .where(PhoneCheckItem.job_id == job_id)
            .order_by(PhoneCheckItem.id)
            .limit(500)
        )
        items = [{
            "id": item.id,
            "account_id": item.account_id,
            "target": item.normalized_phone or item.original_phone,
            "status": item.status,
            "attempts": item.attempts,
            "error_code": item.error_code,
            "error_detail": item.error_detail,
            "started_at": item.started_at,
            "finished_at": item.finished_at,
            "next_retry_at": item.next_retry_at,
        } for item in res.scalars().all()]
    else:
        res = await db.execute(
            select(BulkJobItem)
            .where(BulkJobItem.job_id == job_id)
            .order_by(BulkJobItem.id)
        )
        items = [{
            "id": item.id,
            "account_id": item.account_id,
            "status": item.status,
            "attempts": item.attempts,
            "error_code": item.error_code,
            "error_detail": item.error_detail,
            "started_at": item.started_at,
            "finished_at": item.finished_at,
        } for item in res.scalars().all()]
    out = _job_dict(job)
    out["items"] = items
    counts = Counter(str(item.get("status") or "unknown") for item in items)
    out["status_counts"] = dict(counts)
    timeline = await build_job_timeline(db, job, items)
    out.update(timeline)
    if job.type == "message_multi_send":
        safe_retry = sum(1 for item in items if item["status"] == "pending" and item.get("error_code") == "FloodWaitError" and int(item.get("attempts") or 0) == 0)
        delivered = int(counts.get("ok", 0))
        out["delivery"] = {
            "delivered": delivered,
            "failed": int(counts.get("failed", 0)) + int(counts.get("in_flight_unknown", 0)),
            "pending": sum(int(counts.get(k, 0)) for k in ("queued", "processing", "rate_limited", "pending")),
            "skipped": int(counts.get("skipped", 0)) + int(counts.get("cancelled", 0)),
            "attempted": sum(1 for item in items if int(item.get("attempts") or 0) > 0),
            "safe_retry": safe_retry,
            "delivery_rate": round((delivered / len(items)) * 100, 2) if items else 0.0,
        }
    return out


_EXPORT_FIELDS = [
    "id", "account_id", "target", "status", "attempts", "error_code",
    "error_detail", "started_at", "finished_at", "next_retry_at",
]

def _export_value(value):
    if value is None:
        return ""
    return value.isoformat() if hasattr(value, "isoformat") else str(value)

@router.get("/{job_id}/export.csv")
async def export_job_csv(job_id: str, db: AsyncSession = Depends(get_db)):
    detail = await get_job(job_id, db)
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(_EXPORT_FIELDS)
    for item in detail.get("items", []):
        writer.writerow([_export_value(item.get(field)) for field in _EXPORT_FIELDS])
    data = output.getvalue().encode("utf-8-sig")
    await emit_event("export", "success", "job_csv", "Đã tạo file CSV của tác vụ", job_id=job_id, metadata={"rows": len(detail.get("items", [])), "format": "csv"})
    return StreamingResponse(io.BytesIO(data), media_type="text/csv; charset=utf-8", headers={
        "Content-Disposition": f'attachment; filename="job-{job_id}.csv"'
    })


@router.get("/{job_id}/export.xlsx")
async def export_job_xlsx(job_id: str, db: AsyncSession = Depends(get_db)):
    from openpyxl import Workbook
    detail = await get_job(job_id, db)
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("results")
    ws.append(_EXPORT_FIELDS)
    for item in detail.get("items", []):
        ws.append([_export_value(item.get(field)) for field in _EXPORT_FIELDS])
    summary = wb.create_sheet("summary")
    summary.append(["field", "value"])
    for field in ("id", "type", "status", "total", "success", "failed", "skipped", "pending"):
        summary.append([field, _export_value(detail.get(field))])
    if detail.get("delivery"):
        for key, value in detail["delivery"].items():
            summary.append([f"delivery.{key}", value])
    output = io.BytesIO()
    wb.save(output); output.seek(0)
    await emit_event("export", "success", "job_xlsx", "Đã tạo file XLSX của tác vụ", job_id=job_id, metadata={"rows": len(detail.get("items", [])), "format": "xlsx"})
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={
        "Content-Disposition": f'attachment; filename="job-{job_id}.xlsx"'
    })


@router.post("/{job_id}/pause")
async def pause_job(job_id: str, db: AsyncSession = Depends(get_db)):
    job = await db.get(BulkJob, job_id)
    if not job:
        raise HTTPException(404, "Không tìm thấy tác vụ")
    if job.type not in job_store.RESUMABLE_JOB_TYPES:
        raise HTTPException(400, "Loại tác vụ này không hỗ trợ tạm dừng an toàn")
    accepted = await job_store.pause_job(job_id)
    if not accepted:
        raise HTTPException(409, "Tác vụ hiện không thể tạm dừng")
    await log_audit("job:pause", detail={"job_id": job_id})
    return {"ok": True, "job_id": job_id, "status": "paused"}


@router.post("/{job_id}/resume")
async def resume_job(job_id: str, db: AsyncSession = Depends(get_db)):
    job = await db.get(BulkJob, job_id)
    if not job:
        raise HTTPException(404, "Không tìm thấy tác vụ")
    if job.type not in job_store.RESUMABLE_JOB_TYPES:
        raise HTTPException(400, "Loại tác vụ này không hỗ trợ tiếp tục an toàn")
    accepted = await job_store.resume_job(job_id)
    if not accepted:
        raise HTTPException(409, "Tác vụ hiện không thể tiếp tục")
    await log_audit("job:resume", detail={"job_id": job_id})
    return {"ok": True, "job_id": job_id, "status": "queued"}


@router.post("/{job_id}/cancel")
async def cancel_job(job_id: str, db: AsyncSession = Depends(get_db)):
    job = await db.get(BulkJob, job_id)
    if not job:
        raise HTTPException(404, "Không tìm thấy tác vụ")
    accepted = await job_store.request_cancel(job_id)
    if accepted:
        await log_audit("job:cancel", detail={"job_id": job_id})
    return {"ok": accepted, "job_id": job_id, "status": "cancelling" if accepted else job.status}


class RetryMessageIn(BaseModel):
    text: str = Field(min_length=1, max_length=4096)


@router.post("/{job_id}/retry-message")
async def retry_message_job(job_id: str, body: RetryMessageIn, db: AsyncSession = Depends(get_db)):
    job = await db.get(BulkJob, job_id)
    if not job or job.type != "message_multi_send":
        raise HTTPException(404, "Không tìm thấy tác vụ gửi nhiều người nhận")
    rows = (await db.execute(
        select(MessageDispatchItem).where(
            MessageDispatchItem.job_id == job_id,
            MessageDispatchItem.status == "pending",
            MessageDispatchItem.error_code == "FloodWaitError",
            MessageDispatchItem.attempts == 0,
        ).order_by(MessageDispatchItem.id)
    )).scalars().all()
    if not rows:
        raise HTTPException(409, "Không có người nhận FloodWait chưa gửi nào có thể chạy lại an toàn")
    account_ids = list(dict.fromkeys(int(r.account_id) for r in rows if r.account_id is not None))
    accs = (await db.execute(
        select(Account).where(Account.id.in_(account_ids), Account.deleted_at.is_(None))
    )).scalars().all()
    accounts = [(a.id, a.phone, (f"{a.first_name or ''} {a.last_name or ''}".strip() or a.phone)) for a in accs]
    from ..message_dispatch import eligible_message_accounts, multi_target_message_stream, normalize_message_target
    eligible, excluded = await eligible_message_accounts(accounts)
    if not eligible:
        raise HTTPException(409, "Không có tài khoản sẵn sàng để chạy lại")
    targets = [normalize_message_target(row.target) for row in rows]
    await log_audit("job:retry_message", detail={"source_job_id": job_id, "targets": len(targets)})
    return StreamingResponse(
        multi_target_message_stream(
            eligible, targets, body.text.strip(),
            {"retry_of": job_id, "safe_retry": "pre_send_flood_wait"},
        ),
        media_type="application/x-ndjson",
    )


@router.post("/{job_id}/retry")
async def retry_job(job_id: str, db: AsyncSession = Depends(get_db)):
    job = await db.get(BulkJob, job_id)
    if not job:
        raise HTTPException(404, "Không tìm thấy tác vụ")
    if job.status in {"queued", "running", "cancelling"}:
        raise HTTPException(409, "Không thể chạy lại tác vụ đang hoạt động")
    if job.type not in RETRYABLE_JOB_TYPES:
        raise HTTPException(400, "Loại tác vụ này không thể chạy lại an toàn")

    item_res = await db.execute(
        select(BulkJobItem)
        .where(BulkJobItem.job_id == job_id, BulkJobItem.status.in_(RETRY_ITEM_STATUSES))
        .order_by(BulkJobItem.id)
    )
    items = [item for item in item_res.scalars().all() if item.account_id is not None]
    if not items:
        raise HTTPException(409, "Không có tài khoản lỗi hoặc đang chờ để chạy lại")

    account_ids = [int(item.account_id) for item in items]
    acc_res = await db.execute(
        select(Account).where(Account.id.in_(account_ids), Account.deleted_at.is_(None))
    )
    by_id = {a.id: a for a in acc_res.scalars().all()}
    accounts = [
        (a.id, a.phone, (f"{a.first_name or ''} {a.last_name or ''}".strip() or a.phone))
        for aid in account_ids if (a := by_id.get(aid)) is not None
    ]
    if not accounts:
        raise HTTPException(409, "Không còn tài khoản nào có thể chạy lại")

    params = dict(job.parameters or {})
    params["retry_of"] = job_id
    action = None
    on_success = None

    if job.type == "group_join":
        target = str(params.get("target") or "").strip()
        if not target:
            raise HTTPException(409, "Tác vụ gốc thiếu dữ liệu cần thiết để chạy lại")
        from .groups import _join_handle, _cache
        action = lambda cli, aid: _join_handle(cli, target)
        on_success = lambda aid: _cache.pop(aid, None)

    elif job.type == "group_leave":
        try:
            chat_id = int(params.get("chat_id"))
        except (TypeError, ValueError):
            raise HTTPException(409, "Tác vụ gốc thiếu dữ liệu cần thiết để chạy lại")
        from .groups import _leave_with_client, _cache
        action = lambda cli, aid: _leave_with_client(cli, chat_id)
        on_success = lambda aid: _cache.pop(aid, None)

    elif job.type == "group_leave_target":
        target = str(params.get("target") or "").strip()
        if not target:
            raise HTTPException(409, "Tác vụ gốc thiếu dữ liệu cần thiết để chạy lại")
        from .groups import _leave_by_target_with_client, _cache
        action = lambda cli, aid: _leave_by_target_with_client(cli, target)
        on_success = lambda aid: _cache.pop(aid, None)

    elif job.type == "message_view":
        post_link = str(params.get("post_link") or "").strip()
        if not post_link:
            raise HTTPException(409, "Tác vụ gốc thiếu dữ liệu cần thiết để chạy lại")
        from .messaging import _parse_post_link
        from telethon.tl.functions.messages import GetMessagesViewsRequest
        try:
            chan, msg_id = _parse_post_link(post_link)
        except Exception as exc:
            raise HTTPException(400, friendly_error(exc))

        async def action(cli, aid):
            entity = await cli.get_entity(chan)
            result = await cli(GetMessagesViewsRequest(peer=entity, id=[msg_id], increment=True))
            count = result.views[0].views if (result and result.views) else None
            return "ok", (f"{count} lượt xem" if count is not None else "")

    elif job.type == "terminate_other_sessions":
        from .security import _terminate_other_authorizations

        async def action(cli, aid):
            killed, failed = await _terminate_other_authorizations(cli)
            detail = "không có phiên nào khác" if killed == 0 else f"đã chấm dứt {killed} phiên"
            if failed:
                detail += f", {failed} lỗi"
            return "ok", detail

    await log_audit("job:retry", detail={"source_job_id": job_id, "type": job.type, "accounts": len(accounts)})
    return StreamingResponse(
        bulk_stream(
            accounts, action, on_success=on_success,
            job_type=job.type, job_parameters=params,
        ),
        media_type="application/x-ndjson",
    )
